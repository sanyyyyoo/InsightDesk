"""
Fresh labelling batch for a genuinely untouched final test (test set 2).

Every one of the 338 scored reviews of test set 1 has now been used for
development or held-out evaluation, so the final configuration needs reviews
no method has seen. This batch: 200 NEW Play Store reviews, Zendesk and
Freshdesk only (duplicates are per product, and these two have the most):
  all remaining 1-2 star reviews (147) + 53 random 3-star reviews.
30 of them are labelled by all three people (agreement check); the other 170
are split evenly. Same codebook and single-tab layout as v3, with the three
re-check rules written into the instructions.

Output: labelling/playstore_testset2_sheet.xlsx (new file, local only)

  python experiments/dedup/build_playstore_testset2_sheet.py
"""

import os

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(__file__)
OUT = os.path.join(HERE, "labelling", "playstore_testset2_sheet.xlsx")
KEY_OUT = os.path.join(HERE, "labelling", "playstore_testset2_key.csv")
CODEBOOK = os.path.join(HERE, "labels", "playstore_codebook.csv")
V2_KEY = os.path.join(HERE, "labelling", "playstore_labelling_sheet_v2.xlsx")
LABELLERS = ["Sanyogita", "Krisha", "Krushnam"]
PRODUCTS = ["Zendesk", "Freshdesk"]
N_THREE_STAR = 53
N_OVERLAP = 30
SEED = 26
FIRST_ROW = 451                     # continues the numbering of test set 1 (rows 1-450)
HEADER_ROW = 11
CODE_COL = 11                       # column K
CODE_LAST_ROW = 400

INSTRUCTIONS = [
    "InsightDesk duplicate labelling - TEST SET 2 (new reviews). Pick the SPECIFIC problem. Same product + same code = duplicates.",
    "1. Filter column A (labeller) to your name. Do your Block rows, then the Overlap rows (Overlap: alone, don't look at the others' answers).",
    "2. problem_code: pick from the dropdown. Codes with definitions and example rows are on the right (column K onwards).",
    "RULE 1: a vague rant with no failure ('Very bad', 'Waste app', 'I hate it') -> not_a_problem. Never 'other'.",
    "RULE 2: 'other' only when NO listed code fits - read the code list first. Describe the problem in notes.",
    "RULE 3: *_unspecified only when the review names NO symptom ('can't login', 'keeps crashing', 'doesn't work').",
    "   If it names one (SSO, token expired, white screen, every time I open, too slow), use that specific code.",
    "Two problems -> the main one (usually the one that blocks use); the other code in notes. Unsure -> closest code + confidence = unsure.",
    "About 20-30 s per review, ~85 reviews each. These reviews are new: nobody has tested the pipeline on them.",
]
COLS = ["labeller", "task", "row", "product", "rating", "review_text", "problem_code", "confidence", "notes"]
WIDTHS = {"labeller": 12, "task": 9, "row": 6, "product": 11, "rating": 7, "review_text": 75, "problem_code": 30,
          "confidence": 11, "notes": 28}


def sample():
    d = pd.read_csv("data/play_store_reviews.csv").drop_duplicates("complaint_text")
    used = pd.read_excel(V2_KEY, sheet_name="All_rows_key").complaint_id
    d = d[~d.complaint_id.isin(used) & d.source_product.isin(PRODUCTS)]
    low = d[d.rating <= 2]
    mid = d[d.rating == 3].sample(N_THREE_STAR, random_state=SEED)
    s = pd.concat([low, mid]).sample(frac=1, random_state=SEED).reset_index(drop=True)
    s["row"] = range(FIRST_ROW, FIRST_ROW + len(s))
    rng = np.random.default_rng(SEED)
    overlap = set(rng.choice(len(s), N_OVERLAP, replace=False).tolist())
    s["in_overlap"] = [i in overlap for i in range(len(s))]
    rest = [i for i in range(len(s)) if i not in overlap]
    owner = {i: LABELLERS[k % 3] for k, i in enumerate(rng.permutation(rest))}
    s["block"] = [owner.get(i, "") for i in range(len(s))]
    return s.rename(columns={"source_product": "product", "complaint_text": "review_text"})


def main():
    s = sample()
    parts = []
    for name in LABELLERS:
        for task, rows in (("Block", s[s.block == name]), ("Overlap", s[s.in_overlap])):
            parts.append(rows.sort_values(["product", "row"]).assign(labeller=name, task=task,
                                                                     problem_code="", confidence="", notes=""))
    t = pd.concat(parts, ignore_index=True)[COLS]
    codes = pd.read_csv(CODEBOOK)

    wb = Workbook()
    ws = wb.active
    ws.title = "Labelling"
    blue = PatternFill("solid", fgColor="1F5FA8")
    fill_in = PatternFill("solid", fgColor="FFF7DC")
    thin = Border(top=Side(style="thin", color="1F5FA8"))
    for i, line in enumerate(INSTRUCTIONS, 1):
        ws.cell(i, 1, line).font = Font(bold=(i == 1 or line.startswith("RULE")), size=12 if i == 1 else 10,
                                        color="1F5FA8" if i == 1 else "000000")
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=len(COLS))
    for j, c in enumerate(COLS, 1):
        cell = ws.cell(HEADER_ROW, j, c)
        cell.font, cell.fill = Font(bold=True, color="FFFFFF"), blue
        ws.column_dimensions[get_column_letter(j)].width = WIDTHS[c]
    first, last = HEADER_ROW + 1, HEADER_ROW + len(t)
    prev = None
    for i, rec in enumerate(t.itertuples(index=False), first):
        for j, v in enumerate(rec, 1):
            cell = ws.cell(i, j, v)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if COLS[j - 1] in ("problem_code", "confidence", "notes"):
                cell.fill = fill_in
        if prev is not None and (rec.labeller, rec.task) != prev:
            for j in range(1, len(COLS) + 1):
                ws.cell(i, j).border = thin
        prev = (rec.labeller, rec.task)
    ws.auto_filter.ref = f"A{HEADER_ROW}:{get_column_letter(len(COLS))}{last}"
    ws.freeze_panes = f"A{HEADER_ROW + 1}"

    code_cols = ["problem_code", "area", "definition", "example_rows"]
    for j, c in enumerate(code_cols, CODE_COL):
        cell = ws.cell(HEADER_ROW, j, c)
        cell.font, cell.fill = Font(bold=True, color="FFFFFF"), blue
        ws.column_dimensions[get_column_letter(j)].width = {"definition": 60, "problem_code": 28}.get(c, 16)
    for i, rec in enumerate(codes[code_cols].fillna("").itertuples(index=False), HEADER_ROW + 1):
        for j, v in enumerate(rec, CODE_COL):
            ws.cell(i, j, v).alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions[get_column_letter(CODE_COL - 1)].width = 3
    L = get_column_letter(CODE_COL)
    dv_code = DataValidation(type="list", allow_blank=True, showErrorMessage=True,
                             formula1=f"${L}${HEADER_ROW + 1}:${L}${CODE_LAST_ROW}",   # no leading '=' (Google Sheets)
                             errorTitle="Unknown code",
                             error="Pick a code from the list, or add a new one at the bottom of the code list first.")
    dv_conf = DataValidation(type="list", formula1='"sure,unsure"', allow_blank=True)
    ws.add_data_validation(dv_code)
    ws.add_data_validation(dv_conf)
    pc, cf = (get_column_letter(COLS.index(c) + 1) for c in ("problem_code", "confidence"))
    dv_code.add(f"{pc}{first}:{pc}{last}")
    dv_conf.add(f"{cf}{first}:{cf}{last}")
    wb.save(OUT)
    s[["row", "complaint_id", "product", "rating", "review_text", "created_at", "block", "in_overlap"]].to_csv(
        KEY_OUT, index=False)

    # checks
    used = set(pd.read_excel(V2_KEY, sheet_name="All_rows_key").complaint_id)
    assert not set(s.complaint_id) & used, "overlaps test set 1"
    assert s.complaint_id.is_unique and len(s) == 147 + N_THREE_STAR
    assert (t[t.task == "Overlap"].groupby("row").labeller.nunique() == 3).all()
    assert set(t[t.task == "Block"].row) | set(s[s.in_overlap].row) == set(s.row)
    print(s.groupby(["product", "rating"]).size().unstack(fill_value=0))
    print(t.groupby(["labeller", "task"]).size().unstack(fill_value=0))
    print("reviews:", len(s), "| sheet rows:", len(t), "| checks OK | saved", OUT)


if __name__ == "__main__":
    main()
