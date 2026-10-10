"""
Build the team labelling sheet for real in-domain duplicate labels.

Sample: 300 of the 711 real 1-2 star Play Store reviews
(data/play_store_reviews.csv), stratified by product, fixed seed.
Split into 3 blocks of 100 (one per labeller) plus a 50-review overlap that
all three label independently (for inter-annotator agreement).

Suggested groups are only a starting point to speed labelling: batch
average-linkage over MiniLM cosine of the boilerplate-suppressed text,
cut STRICTER (0.55) than the v2 operating point, because v2 at 0.44
over-merged real reviews. Rows are sorted by suggested group so similar
reviews sit together. Labellers must correct them; the final labels are the
human ones.

Output: experiments/dedup/labelling/playstore_labelling_sheet.xlsx (local only)

  python experiments/dedup/build_playstore_labelling.py
"""

import os
import sys

import numpy as np
import pandas as pd

sys.path.append(os.path.dirname(__file__))
sys.path.append(os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..")))

from backend.services.preprocess import clean_text  # noqa: E402
from compare_pipelines import batch_cos  # noqa: E402
from normalize import boilerplate_suppress  # noqa: E402

SEED = 11
N_SAMPLE = 300
N_OVERLAP = 50
SUGGEST_T = 0.55
LABELLERS = ["Sanyogita", "Krisha", "Krushnam"]
OUT = os.path.join(os.path.dirname(__file__), "labelling", "playstore_labelling_sheet.xlsx")

INSTRUCTIONS = [
    ("Goal", "Give every review a PROBLEM LABEL. Reviews with the same label = the same case (duplicates). "
             "These human labels become the real-data test set for InsightDesk dedup."),
    ("Your block", "Label the 100 rows in your own sheet (Block_<name>), then ALL 50 rows in Overlap_<name>. "
                   "Do the overlap on your own - do not look at each other's answers (we measure agreement)."),
    ("Same case rule", "Same case = same product area AND same specific failure, i.e. one fix would solve both "
                       "(Krisha's rule from ground_truth_corrections.csv). Different wording does not matter; "
                       "different feature or different symptom = different case."),
    ("Examples: SAME", "'code never arrives on email' + 'not getting passcode on email' -> login_code_not_received. "
                       "'app keeps crashing' + 'crashes every time I open it' -> app_crashes."),
    ("Examples: DIFFERENT", "'notifications arrive late' vs 'can't delete notifications' vs 'no notifications at all' "
                            "-> three labels. 'wouldn't stop asking for a review' vs 'manager review ignored' -> different "
                            "(shared word only)."),
    ("is_problem_report", "Y = reports a problem with the product. N = praise, feature request with no failure, "
                          "pricing opinion, rant with no specific problem. For N rows use label not_a_problem."),
    ("problem_label", "Short snake_case: <area>_<failure>, e.g. login_sso_fails, tickets_not_loading, "
                      "notifications_delayed. Reuse labels from the Labels sheet; add a new one there with a one-line "
                      "definition when nothing fits. A review with two problems: use the main one, note the other."),
    ("confidence", "sure / unsure. Mark unsure instead of guessing - unsure rows are discussed together."),
    ("suggested_group", "A machine suggestion (similar-text cluster) to speed you up. It is often WRONG - "
                        "rows are sorted by it only so similar reviews sit together. Never copy it blindly."),
    ("Time", "About 1-1.5 hours per person (100 + 50 short reviews)."),
    ("Privacy", "Public Play Store review text; do not add names, emails or anything identifying in notes."),
]


def main():
    from sentence_transformers import SentenceTransformer
    d = pd.read_csv("data/play_store_reviews.csv")
    d = d[d.rating <= 2].drop_duplicates("complaint_text").reset_index(drop=True)
    rng = np.random.default_rng(SEED)
    frac = N_SAMPLE / len(d)
    parts = [g.sample(n=max(1, round(len(g) * frac)), random_state=SEED) for _, g in d.groupby("source_product")]
    s = pd.concat(parts)
    s = s.sample(n=N_SAMPLE, random_state=SEED) if len(s) > N_SAMPLE else s
    s = s.reset_index(drop=True)

    enc = SentenceTransformer("sentence-transformers/all-MiniLM-L6-v2")
    text = pd.Series(boilerplate_suppress(s.complaint_text.astype(str).tolist(), enc)).map(clean_text)
    emb = enc.encode(text.tolist(), normalize_embeddings=True, show_progress_bar=False)
    groups = batch_cos(emb, SUGGEST_T)
    sizes = pd.Series(groups).value_counts()
    s["suggested_group"] = [f"G{g:03d}" if sizes[g] > 1 else "single" for g in groups]

    s["row"] = range(1, len(s) + 1)
    s["block"] = [LABELLERS[i % 3] for i in rng.permutation(len(s))]
    overlap_rows = set(rng.choice(len(s), size=N_OVERLAP, replace=False).tolist())
    s["in_overlap"] = [i in overlap_rows for i in range(len(s))]

    base_cols = ["row", "complaint_id", "source_product", "rating", "complaint_text", "suggested_group"]
    fill_cols = ["is_problem_report", "problem_label", "confidence", "notes"]

    def sheet(frame):
        f = frame.sort_values(["suggested_group", "source_product", "row"])[base_cols].copy()
        for c in fill_cols:
            f[c] = ""
        return f.rename(columns={"source_product": "product", "complaint_text": "review_text"})

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    with pd.ExcelWriter(OUT, engine="openpyxl") as xw:
        pd.DataFrame(INSTRUCTIONS, columns=["Topic", "Instruction"]).to_excel(xw, sheet_name="Instructions", index=False)
        for name in LABELLERS:
            sheet(s[s.block == name]).to_excel(xw, sheet_name=f"Block_{name}", index=False)
        for name in LABELLERS:
            sheet(s[s.in_overlap]).to_excel(xw, sheet_name=f"Overlap_{name}", index=False)
        pd.DataFrame({"problem_label": ["not_a_problem"],
                      "definition": ["No specific product failure (praise, feature wish, pricing opinion, rant)"],
                      "added_by": ["template"]}).to_excel(xw, sheet_name="Labels", index=False)
        s[base_cols + ["block", "in_overlap"]].to_excel(xw, sheet_name="All_rows_key", index=False)

        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.worksheet.datavalidation import DataValidation
        wb = xw.book
        for ws in wb.worksheets:
            ws.freeze_panes = "A2"
            for c in ws[1]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor="1F5FA8")
            widths = {"Instruction": 110, "review_text": 80, "Topic": 20, "notes": 30, "problem_label": 28,
                      "definition": 70}
            for col in ws.iter_cols(min_row=1, max_row=1):
                name = col[0].value
                ws.column_dimensions[col[0].column_letter].width = widths.get(name, 14)
            for row in ws.iter_rows(min_row=2):
                for c in row:
                    c.alignment = Alignment(wrap_text=True, vertical="top")
            if ws.title.startswith(("Block_", "Overlap_")):
                n = ws.max_row
                yn = DataValidation(type="list", formula1='"Y,N"', allow_blank=True)
                conf = DataValidation(type="list", formula1='"sure,unsure"', allow_blank=True)
                ws.add_data_validation(yn)
                ws.add_data_validation(conf)
                yn.add(f"G2:G{n}")
                conf.add(f"I2:I{n}")
                for r in range(2, n + 1):
                    for col in "GHIJ":
                        ws[f"{col}{r}"].fill = PatternFill("solid", fgColor="FFF7DC")

    print(f"sample {len(s)} reviews: {s.source_product.value_counts().to_dict()}")
    print(f"blocks: {s.block.value_counts().to_dict()} | overlap {s.in_overlap.sum()}")
    print(f"suggested groups (cosine >= {SUGGEST_T}): {int((sizes > 1).sum())} groups covering "
          f"{int(sizes[sizes > 1].sum())} reviews, {int((sizes == 1).sum())} singles")
    print("saved", OUT)


if __name__ == "__main__":
    main()
