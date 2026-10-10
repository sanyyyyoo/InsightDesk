"""
Labelling sheet v3: everything in ONE tab (instructions + all rows + code list).

Same rows and codebook as v2, with Krushnam's 50 Extra reviews moved to
Sanyogita to balance the remaining work. The 50 overlap reviews stay with all
three people (needed for the agreement check).

Layout of the single 'Labelling' tab:
  rows 1-9      instructions
  row 11        table header (frozen); rows below = one row per (labeller, review)
  columns A-L   the labelling table (filter column A to your name)
  columns N-Q   the problem-code list; the dropdown reads column N, so new
                codes typed at the bottom of the list appear in it

Output: labelling/playstore_labelling_sheet_v3.xlsx (new file, local only)

  python experiments/dedup/build_playstore_labelling_v3.py
"""

import os

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from build_playstore_labelling_v2 import CODEBOOK, LABELLERS, extra_rows, load_v1, v2_tab

HERE = os.path.dirname(__file__)
OUT = os.path.join(HERE, "labelling", "playstore_labelling_sheet_v3.xlsx")
MOVE_EXTRA = {"Krushnam": "Sanyogita"}
HEADER_ROW = 11
CODE_COL = 14                       # column N
CODE_LAST_ROW = 400                 # room for new codes

INSTRUCTIONS = [
    "InsightDesk duplicate labelling - pick the SPECIFIC problem for each review. Same product + same code = duplicates.",
    "1. Filter column A (labeller) to your name. Do the rows in order: Block, then Overlap, then Extra.",
    "2. problem_code: pick from the dropdown. The code list with definitions and example rows is on the right (column N onwards).",
    "3. Vague review ('can't login', 'keeps crashing', 'doesn't work' with no detail) -> the *_unspecified code. Do not guess.",
    "4. No failure (praise, question, rant) -> not_a_problem. About another company -> not_about_this_app. Two problems -> main one, other in notes.",
    "5. Nothing fits -> 'other' + describe in notes. Seen it 2-3 times? Add a new code at the bottom of the code list (it joins the dropdown).",
    "6. Overlap rows: do them ALONE and do not look at the other two people's Overlap rows - we measure agreement on them.",
    "7. Unsure -> pick the closest code and set confidence = unsure. Grey old_* columns are your earlier answers, for reference only.",
    "About 20-30 s per review. Work left: Sanyogita 150 re-checks + 100 new; Krisha ~152; Krushnam 150.",
]
TASK_ORDER = {"Block": 0, "Overlap": 1, "Extra": 2}
COLS = ["labeller", "task", "row", "product", "rating", "review_text", "problem_code", "confidence", "notes",
        "old_label", "old_is_problem", "old_notes"]
WIDTHS = {"labeller": 12, "task": 9, "row": 6, "product": 11, "rating": 7, "review_text": 70, "problem_code": 30,
          "confidence": 11, "notes": 26, "old_label": 20, "old_is_problem": 9, "old_notes": 16}


def build_table():
    v1, key = load_v1()
    extra = extra_rows(key)
    extra["block"] = extra.block.replace(MOVE_EXTRA)
    rows = key.rename(columns={"source_product": "product", "complaint_text": "review_text"})
    parts = []
    for name in LABELLERS:
        parts.append(v2_tab(rows[rows.block == name], v1[f"Block_{name}"]).assign(labeller=name, task="Block"))
        parts.append(v2_tab(rows[rows.in_overlap], v1[f"Overlap_{name}"]).assign(labeller=name, task="Overlap"))
        e = extra[extra.block == name]
        if len(e):
            parts.append(v2_tab(e).assign(labeller=name, task="Extra"))
    t = pd.concat(parts, ignore_index=True)       # each part keeps its own within-task order
    for c in ["old_label", "old_is_problem", "old_notes"]:
        t[c] = t[c].where(t[c].notna(), "")
    return t[COLS], v1, extra


def main():
    t, v1, extra = build_table()
    codes = pd.read_csv(CODEBOOK)
    wb = Workbook()
    ws = wb.active
    ws.title = "Labelling"
    blue = PatternFill("solid", fgColor="1F5FA8")
    fill_in = PatternFill("solid", fgColor="FFF7DC")
    thin = Border(top=Side(style="thin", color="1F5FA8"))

    # instructions
    for i, line in enumerate(INSTRUCTIONS, 1):
        ws.cell(i, 1, line).font = Font(bold=(i == 1), size=12 if i == 1 else 10,
                                        color="1F5FA8" if i == 1 else "000000")
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=len(COLS))

    # table
    for j, c in enumerate(COLS, 1):
        cell = ws.cell(HEADER_ROW, j, c)
        cell.font, cell.fill = Font(bold=True, color="FFFFFF"), blue
        ws.column_dimensions[get_column_letter(j)].width = WIDTHS[c]
    first, last = HEADER_ROW + 1, HEADER_ROW + len(t)
    prev = None
    for i, rec in enumerate(t.itertuples(index=False), first):
        for j, v in enumerate(rec, 1):
            cell = ws.cell(i, j, v)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            name = COLS[j - 1]
            if name in ("problem_code", "confidence", "notes"):
                cell.fill = fill_in
            elif name.startswith("old_"):
                cell.font = Font(color="808080")
        if prev is not None and (rec.labeller, rec.task) != prev:     # line between sections
            for j in range(1, len(COLS) + 1):
                ws.cell(i, j).border = thin
        prev = (rec.labeller, rec.task)
    ws.auto_filter.ref = f"A{HEADER_ROW}:{get_column_letter(len(COLS))}{last}"
    ws.freeze_panes = f"A{HEADER_ROW + 1}"

    # code list on the right
    code_cols = ["problem_code", "area", "definition", "example_rows"]
    for j, c in enumerate(code_cols, CODE_COL):
        cell = ws.cell(HEADER_ROW, j, c)
        cell.font, cell.fill = Font(bold=True, color="FFFFFF"), blue
        ws.column_dimensions[get_column_letter(j)].width = {"definition": 60, "problem_code": 28}.get(c, 16)
    for i, rec in enumerate(codes[code_cols].fillna("").itertuples(index=False), HEADER_ROW + 1):
        for j, v in enumerate(rec, CODE_COL):
            ws.cell(i, j, v).alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions[get_column_letter(CODE_COL - 1)].width = 3

    code_letter = get_column_letter(CODE_COL)
    dv_code = DataValidation(type="list", allow_blank=True, showErrorMessage=True,
                             formula1=f"${code_letter}${HEADER_ROW + 1}:${code_letter}${CODE_LAST_ROW}",
                             errorTitle="Unknown code",
                             error="Pick a code from the list, or add a new one at the bottom of the code list first.")
    dv_conf = DataValidation(type="list", formula1='"sure,unsure"', allow_blank=True)
    ws.add_data_validation(dv_code)
    ws.add_data_validation(dv_conf)
    pc, cf = (get_column_letter(COLS.index(c) + 1) for c in ("problem_code", "confidence"))
    dv_code.add(f"{pc}{first}:{pc}{last}")
    dv_conf.add(f"{cf}{first}:{cf}{last}")
    wb.save(OUT)

    # checks: every earlier answer carried over; each review assigned as intended
    for name in LABELLERS:
        for task in ("Block", "Overlap"):
            a = v1[f"{task}_{name}"].set_index("row").problem_label.dropna().astype(str).sort_index()
            s = t[(t.labeller == name) & (t.task == task)]
            b = s.set_index("row").old_label.astype(str)
            b = b[b != ""].sort_index()
            assert a.equals(b), (name, task)
            assert set(s.row) == set(v1[f"{task}_{name}"].row), (name, task)
    assert set(t[t.task == "Extra"].row) == set(extra.row) and t[t.task == "Extra"].row.is_unique
    assert (t[t.task == "Overlap"].groupby("row").labeller.nunique() == 3).all()
    print(t.groupby(["labeller", "task"]).size().unstack(fill_value=0))
    print("rows:", len(t), "| codes:", len(codes), "| checks OK | saved", OUT)


if __name__ == "__main__":
    main()
