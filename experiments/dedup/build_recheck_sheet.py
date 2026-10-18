"""
Re-check sheet after the first agreement check (Fleiss kappa 0.568 < 0.6).

Flags rows only by rule, never suggests a code (labels stay human):
  blank            - no code yet
  other            - check a listed code really does not fit
  unspecified      - check the review really gives no detail
  overlap_disagree - the three people picked different codes (discussed together)

One tab 'Recheck': instructions on top, rows below (filter by labeller),
code list on the right feeding the new_code dropdown.

Input : labelling/team_labels_v3_table.csv (team's filled sheet, flattened)
Output: labelling/recheck_sheet.xlsx

  python experiments/dedup/build_recheck_sheet.py
"""

import os

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(__file__)
TABLE = os.path.join(HERE, "labelling", "team_labels_v3_table.csv")
CODEBOOK = os.path.join(HERE, "labels", "playstore_codebook.csv")
OUT = os.path.join(HERE, "labelling", "recheck_sheet.xlsx")
HEADER_ROW = 12
CODE_COL = 12                      # column L

INSTRUCTIONS = [
    "Re-check - about 20 minutes each. Agreement on the overlap was kappa 0.568 (target 0.6); most gaps come from 3 rules read differently.",
    "RULE 1: a vague rant with no failure ('Very bad', 'I don't like it', 'Waste app') -> not_a_problem. Never 'other'.",
    "RULE 2: 'other' only when NO listed code fits - read the code list on the right (column L onwards) before keeping 'other'.",
    "RULE 3: *_unspecified only when the review gives NO detail. If it names a symptom (SSO, white screen, token expired, every time I open) use that code.",
    "Part A (your rows): filter column A to your name. If your code is still right, leave new_code empty. Otherwise pick the new code.",
    "Part B (labeller = TEAM): overlap rows where you disagreed. Decide together; put the agreed code in new_code (2-of-3 rows: majority unless you agree otherwise).",
    "The 'flag' column only says WHY a row is listed - it does not mean your answer is wrong.",
    "When done, upload the file back to Drive; I will merge new_code into the labels and recompute kappa (before and after are both reported).",
]
COLS = ["labeller", "flag", "task", "row", "product", "review_text", "current_code", "new_code", "notes"]
WIDTHS = {"labeller": 11, "flag": 15, "task": 9, "row": 6, "product": 10, "review_text": 70,
          "current_code": 34, "new_code": 30, "notes": 22}


def flag_of(code):
    if code == "":
        return "blank"
    if code == "other":
        return "other"
    if code.endswith("unspecified"):
        return "unspecified"
    return None


def main():
    t = pd.read_csv(TABLE, keep_default_na=False)
    t["flag"] = t.problem_code.map(flag_of)
    own = t[t.flag.notna()].assign(current_code=t.problem_code)
    own = own.sort_values(["labeller", "flag", "task", "row"])

    ov = t[t.task == "Overlap"].pivot(index="row", columns="labeller", values="problem_code")
    dis = ov[ov.nunique(axis=1) > 1]
    info = t[t.task == "Overlap"].drop_duplicates("row").set_index("row")
    team = pd.DataFrame({
        "labeller": "TEAM", "flag": "overlap_disagree", "task": "Overlap", "row": dis.index,
        "product": info["product"].reindex(dis.index).values,
        "review_text": info.review_text.reindex(dis.index).values,
        "current_code": [" | ".join(f"{p[:3]}: {c or '(blank)'}" for p, c in r.items()) for _, r in dis.iterrows()],
    })
    rows = pd.concat([own, team], ignore_index=True)
    rows["new_code"] = ""
    rows["notes"] = rows["notes"].fillna("")
    rows = rows[COLS]

    codes = pd.read_csv(CODEBOOK)
    wb = Workbook()
    ws = wb.active
    ws.title = "Recheck"
    blue = PatternFill("solid", fgColor="1F5FA8")
    for i, line in enumerate(INSTRUCTIONS, 1):
        ws.cell(i, 1, line).font = Font(bold=(i == 1 or line.startswith("RULE")), size=11 if i == 1 else 10,
                                        color="1F5FA8" if i == 1 else "000000")
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=len(COLS))
    for j, c in enumerate(COLS, 1):
        cell = ws.cell(HEADER_ROW, j, c)
        cell.font, cell.fill = Font(bold=True, color="FFFFFF"), blue
        ws.column_dimensions[get_column_letter(j)].width = WIDTHS[c]
    first, last = HEADER_ROW + 1, HEADER_ROW + len(rows)
    for i, rec in enumerate(rows.itertuples(index=False), first):
        for j, v in enumerate(rec, 1):
            cell = ws.cell(i, j, v)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if COLS[j - 1] in ("new_code", "notes"):
                cell.fill = PatternFill("solid", fgColor="FFF7DC")
    ws.auto_filter.ref = f"A{HEADER_ROW}:{get_column_letter(len(COLS))}{last}"
    ws.freeze_panes = f"A{HEADER_ROW + 1}"

    code_cols = ["problem_code", "area", "definition", "example_rows"]
    for j, c in enumerate(code_cols, CODE_COL):
        cell = ws.cell(HEADER_ROW, j, c)
        cell.font, cell.fill = Font(bold=True, color="FFFFFF"), blue
        ws.column_dimensions[get_column_letter(j)].width = {"definition": 60, "problem_code": 28}.get(c, 16)
    for i, rec in enumerate(codes[code_cols].fillna("").itertuples(index=False), HEADER_ROW + 1):
        for j, v in enumerate(rec, CODE_COL):
            ws.cell(i, j, v).alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions[get_column_letter(CODE_COL - 1)].width = 3

    L = get_column_letter(CODE_COL)
    dv = DataValidation(type="list", allow_blank=True, showErrorMessage=True,
                        formula1=f"${L}${HEADER_ROW + 1}:${L}${HEADER_ROW + len(codes)}",   # no leading '=' (Google Sheets)
                        errorTitle="Unknown code", error="Pick a code from the list.")
    ws.add_data_validation(dv)
    nc = get_column_letter(COLS.index("new_code") + 1)
    dv.add(f"{nc}{first}:{nc}{last}")
    wb.save(OUT)

    print(rows.groupby(["labeller", "flag"]).size().unstack(fill_value=0))
    print("rows:", len(rows), "| saved", OUT)


if __name__ == "__main__":
    main()
