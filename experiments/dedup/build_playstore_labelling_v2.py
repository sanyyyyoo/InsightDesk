"""
Labelling sheet v2: predefined problem codes + 150 extra reviews.

Why v2: in v1 everyone invented labels, and they came out topic-level
('login_issue' covered five different failures) with spelling variants,
so labels could not be compared across people and Fleiss kappa could not be
computed. v2 gives one shared codebook of SPECIFIC failures
(labels/playstore_codebook.csv, drafted by reading the review texts, not from
pipeline clusters), chosen from a dropdown.

Extra rows: 150 more 1-2 star reviews (100 Zendesk, 50 Freshdesk) not in the
first 300. Duplicates are per product, and these two products have the most
reviews, so they add the most duplicate pairs per labelled row.

Nothing already labelled is lost: every row keeps its v1 answers in the
old_label / old_is_problem columns.

Input : labelling/team_labels_partial.xlsx (team's v1 sheet, downloaded)
Output: labelling/playstore_labelling_sheet_v2.xlsx (new file, local only)

  python experiments/dedup/build_playstore_labelling_v2.py
"""

import os

import numpy as np
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(__file__)
V1 = os.path.join(HERE, "labelling", "team_labels_partial.xlsx")
CODEBOOK = os.path.join(HERE, "labels", "playstore_codebook.csv")
OUT = os.path.join(HERE, "labelling", "playstore_labelling_sheet_v2.xlsx")
LABELLERS = ["Sanyogita", "Krisha", "Krushnam"]
SEED = 11
EXTRA = {"Zendesk": 100, "Freshdesk": 50}
CODE_ROWS = 300          # dropdown range on the Labels sheet, leaves room for new codes

INSTRUCTIONS = [
    ("What changed (v2)", "Labels are now picked from a dropdown of SPECIFIC problems (Labels tab). "
                          "Your v1 answers are kept in old_label / old_is_problem - nothing is lost."),
    ("Rule", "Pick the code that names the specific failure. Same product + same code = same case "
             "(duplicates). One fix would solve both = same code."),
    ("Your work", "Block_<name> (100), Overlap_<name> (50, do it alone - we measure agreement), "
                  "Extra_<name> (50 new reviews)."),
    ("Not sure which code?", "Read the definition and example rows in the Labels tab. Still unsure: pick the "
                             "closest and set confidence = unsure."),
    ("Nothing fits", "Pick 'other' and describe the problem in notes. If you see the same new problem 2-3 "
                     "times, add a new code at the bottom of the Labels tab (it appears in the dropdown) "
                     "and tell the team."),
    ("Vague reviews", "'Can't login' / 'keeps crashing' / 'doesn't work' with no detail -> the *_unspecified "
                      "codes. They are NOT scored as duplicates, so do not guess a specific code."),
    ("No problem", "Praise, questions, rants with no failure -> not_a_problem. About another company -> "
                   "not_about_this_app."),
    ("Two problems", "Pick the main one (usually the first or the one that blocks use); put the other code in notes."),
    ("Time", "About 20-30 s per review with the dropdown."),
]


def load_v1():
    x = pd.read_excel(V1, sheet_name=None)
    return x, x["All_rows_key"]


def extra_rows(key):
    d = pd.read_csv("data/play_store_reviews.csv")
    d = d[d.rating <= 2].drop_duplicates("complaint_text")
    d = d[~d.complaint_id.isin(key.complaint_id)]
    parts = [d[d.source_product == p].sample(n, random_state=SEED) for p, n in EXTRA.items()]
    e = pd.concat(parts).reset_index(drop=True)
    e["row"] = range(301, 301 + len(e))
    rng = np.random.default_rng(SEED)
    e["block"] = [LABELLERS[i % 3] for i in rng.permutation(len(e))]
    return e.rename(columns={"source_product": "product", "complaint_text": "review_text"})


def v2_tab(frame, old=None):
    f = frame[["row", "complaint_id", "product", "rating", "review_text"]].copy()
    if old is not None:
        o = old.set_index("row")
        f["old_label"] = f.row.map(o.problem_label)
        f["old_is_problem"] = f.row.map(o.is_problem_report)
        f["old_notes"] = f.row.map(o.notes)
        # group the rows each person already labelled so one area is checked at a time
        f["_k"] = f.old_label.fillna("~").astype(str).str.lower().str.replace(r"[\s_]+", "", regex=True)
        f = f.sort_values(["product", "_k", "row"]).drop(columns="_k")
    else:
        f = f.sort_values(["product", "row"])
    for c in ["problem_code", "confidence", "notes"]:
        f[c] = ""
    return f


def style(wb, n_codes):
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F5FA8")
        widths = {"Instruction": 110, "review_text": 75, "Topic": 22, "notes": 30, "old_notes": 20,
                  "problem_code": 30, "old_label": 22, "definition": 80, "example_rows": 26, "area": 20}
        for col in ws.iter_cols(min_row=1, max_row=1):
            ws.column_dimensions[col[0].column_letter].width = widths.get(col[0].value, 12)
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        if ws.title.split("_")[0] in ("Block", "Overlap", "Extra"):
            head = {c.value: c.column_letter for c in ws[1]}
            n = ws.max_row
            code = DataValidation(type="list", formula1=f"Labels!$A$2:$A${CODE_ROWS + 1}", allow_blank=True,
                                  showErrorMessage=True, errorTitle="Unknown code",
                                  error="Pick a code from the list, or add a new one in the Labels tab first.")
            conf = DataValidation(type="list", formula1='"sure,unsure"', allow_blank=True)
            ws.add_data_validation(code)
            ws.add_data_validation(conf)
            code.add(f"{head['problem_code']}2:{head['problem_code']}{n}")
            conf.add(f"{head['confidence']}2:{head['confidence']}{n}")
            for r in range(2, n + 1):
                for name in ("problem_code", "confidence", "notes"):
                    ws[f"{head[name]}{r}"].fill = PatternFill("solid", fgColor="FFF7DC")
                for name in ("old_label", "old_is_problem", "old_notes"):
                    if name in head:
                        ws[f"{head[name]}{r}"].font = Font(color="808080")


def main():
    v1, key = load_v1()
    extra = extra_rows(key)
    codes = pd.read_csv(CODEBOOK)
    assert codes.problem_code.is_unique
    rows = key.rename(columns={"source_product": "product", "complaint_text": "review_text"})

    with pd.ExcelWriter(OUT, engine="openpyxl") as xw:
        pd.DataFrame(INSTRUCTIONS, columns=["Topic", "Instruction"]).to_excel(xw, sheet_name="Instructions", index=False)
        codes.to_excel(xw, sheet_name="Labels", index=False)
        for name in LABELLERS:
            v2_tab(rows[rows.block == name], v1[f"Block_{name}"]).to_excel(xw, sheet_name=f"Block_{name}", index=False)
        for name in LABELLERS:
            v2_tab(rows[rows.in_overlap], v1[f"Overlap_{name}"]).to_excel(xw, sheet_name=f"Overlap_{name}", index=False)
        for name in LABELLERS:
            v2_tab(extra[extra.block == name]).to_excel(xw, sheet_name=f"Extra_{name}", index=False)
        allk = pd.concat([rows[["row", "complaint_id", "product", "rating", "review_text", "block", "in_overlap"]],
                          extra[["row", "complaint_id", "product", "rating", "review_text", "block"]]
                          .assign(in_overlap=False)])
        allk.to_excel(xw, sheet_name="All_rows_key", index=False)
        style(xw.book, len(codes))

    # checks: every v1 answer carried over, row sets unchanged, extras new
    out = pd.read_excel(OUT, sheet_name=None)
    for t in [f"{k}_{n}" for k in ("Block", "Overlap") for n in LABELLERS]:
        a = v1[t].set_index("row").problem_label.dropna().astype(str)
        b = out[t].set_index("row").old_label.dropna().astype(str)
        assert a.sort_index().equals(b.sort_index()), t
        assert set(v1[t].row) == set(out[t].row), t
    assert not set(extra.complaint_id) & set(key.complaint_id)
    print("codes:", len(codes), "| extra:", extra["product"].value_counts().to_dict(),
          "| per person:", extra.block.value_counts().to_dict())
    print("v1 answers carried over: OK | saved", OUT)


if __name__ == "__main__":
    main()
