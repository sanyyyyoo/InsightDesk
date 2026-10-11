"""
Blind A/B sheet for the RCA v2 test (rule: rca_v2_rule.md).

Each of the 20 real-review cases shows its complaints and two RCAs, "RCA A" and
"RCA B"; which one is v1 and which v2 is random per case (seed fixed) and kept
in rca_ab_key.csv, NOT in the sheet.

  python reports/replay/build_rca_ab_sheet.py                       # v1 vs v2, top 20 (round 2)
  python reports/replay/build_rca_ab_sheet.py next20_real_case_review_v1.xlsx          next20_real_case_review_v3.xlsx v3 rca_ab3                 # v1 vs v3, ranks 21-40
"""

import os
import random
import sys

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
OLD, NEW, NEWV, PREFIX = (sys.argv[1:5] if len(sys.argv) >= 5 else
                          ("top20_real_case_review.xlsx", "top20_real_case_review_v2.xlsx", "v2", "rca_ab"))
v1 = pd.read_excel(os.path.join(HERE, OLD))
v2 = pd.read_excel(os.path.join(HERE, NEW)).set_index("case_id")
assert set(v1.case_id) == set(v2.index), "v1 and v2 sheets must cover the same cases"

rng = random.Random(20261011)
rows, key = [], []
for n, r in enumerate(v1.itertuples(), 1):
    old = (r.rca_probable_cause, r.rca_recommended_action)
    new = (v2.at[r.case_id, "rca_probable_cause"], v2.at[r.case_id, "rca_recommended_action"])
    a_is_v2 = rng.random() < 0.5
    A, B = (new, old) if a_is_v2 else (old, new)
    samples = [getattr(r, f"sample_{k}", "") if isinstance(getattr(r, f"sample_{k}", ""), str) else "" for k in range(1, 7)]
    rows.append([n, r.complaints, r.products] + samples + [A[0], A[1], B[0], B[1]])
    key.append({"case": n, "case_id": r.case_id, "A": NEWV if a_is_v2 else "v1", "B": "v1" if a_is_v2 else NEWV})

cols = ["case", "complaints_in_case", "company"] + [f"complaint_{k}" for k in range(1, 7)] + \
       ["RCA_A_cause", "RCA_A_action", "RCA_B_cause", "RCA_B_action",
        "A_useful", "A_unsupported_claim", "B_useful", "B_unsupported_claim", "better", "notes"]
wb = Workbook()
ins = wb.active
ins.title = "Instructions"
for k, t in enumerate([
        "Each row is one complaint case with TWO root-cause analyses, RCA A and RCA B. Rate both, on your own,",
        "judging only what you see (the complaints and the RCA text).",
        "",
        "A_useful / B_useful - would this probable cause + recommended action help the team that has to fix it?",
        "  yes = specific, consistent with the complaints, a concrete next step",
        "  partly = plausible but generic, or only part fits the complaints",
        "  no = wrong, contradicts the complaints, or nothing to act on",
        "A_unsupported_claim / B_unsupported_claim - does the RCA state something the complaints do NOT give",
        "  (a cause presented as certain, a version / platform / component / update not mentioned)?  yes / no",
        "better - which RCA would you rather give the team: A / B / same",
        "notes - one short line on the main difference."], 1):
    ins.cell(k, 1, t)
ins.column_dimensions["A"].width = 110
ws = wb.create_sheet("Cases")
blue, fill_in = PatternFill("solid", fgColor="1F5FA8"), PatternFill("solid", fgColor="FFF7DC")
for j, c in enumerate(cols, 1):
    cell = ws.cell(1, j, c)
    cell.font, cell.fill = Font(bold=True, color="FFFFFF"), blue
    ws.column_dimensions[get_column_letter(j)].width = 45 if c.startswith(("complaint_", "RCA_")) else 14
for i, row in enumerate(rows, 2):
    for j, v in enumerate(row + [""] * 6, 1):
        c = ws.cell(i, j, v)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if j > len(cols) - 6:
            c.fill = fill_in
for col, opts in [("A_useful", "yes,partly,no"), ("B_useful", "yes,partly,no"), ("A_unsupported_claim", "yes,no"),
                  ("B_unsupported_claim", "yes,no"), ("better", "A,B,same")]:
    L = get_column_letter(cols.index(col) + 1)
    dv = DataValidation(type="list", formula1=f'"{opts}"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{L}2:{L}{len(rows) + 1}")
ws.freeze_panes = "D2"
wb.save(os.path.join(HERE, f"{PREFIX}_sheet.xlsx"))
pd.DataFrame(key).to_csv(os.path.join(HERE, f"{PREFIX}_key.csv"), index=False)
open(os.path.join(HERE, f"{PREFIX}_prompt.txt"), "w", encoding="utf-8").write(
    """The attached spreadsheet has 20 complaint cases from app-store reviews of customer-support apps. Each row shows up to 6 complaints and TWO root-cause analyses written for that case, "RCA A" and "RCA B". The "Instructions" tab defines the ratings.

Rate every case yourself by reading it, the way a human reviewer would. Do not write code or use keyword rules.

For each of the 20 rows fill:
- A_useful and B_useful: yes / partly / no
- A_unsupported_claim and B_unsupported_claim: yes / no (does the RCA state something the complaints do not give: a cause presented as certain, or a version / platform / component / update not mentioned?)
- better: A / B / same
- notes: one short line on the main difference

Finally, in your reply, give an overall quality score out of 10 for the A RCAs as a set and for the B RCAs as a set, with one sentence each.

Keep the case column unchanged and give me the completed spreadsheet as an .xlsx file to download.""")
print(f"sheet with {len(rows)} cases; A is v2 in {sum(k['A'] == NEWV for k in key)} of them (key in {PREFIX}_key.csv)")
