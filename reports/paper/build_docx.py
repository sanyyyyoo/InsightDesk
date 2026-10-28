"""Build the IEEE-style two-column Word version of the paper from the doc's XML (paper.xml).

Figures are the doc's own renders (img/fig*.png); equations are rendered with matplotlib mathtext.
  python build_docx.py   ->  InsightDesk_paper.docx
"""
import html
import re

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from docx import Document  # noqa: E402
from docx.enum.section import WD_SECTION  # noqa: E402
from docx.enum.table import WD_TABLE_ALIGNMENT  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Inches, Pt  # noqa: E402

FONT = "Times New Roman"
COL_W = Inches(3.45)

FIGS = {"04f74690-b065": ("img/fig1.png", 5.2, "Flowchart of the InsightDesk pipeline (t_low = 0.40 for app reviews and chat, 0.26 for tickets and bug reports)"),
        "ff082679-c090": ("img/fig2.png", 7.0, "System architecture: nine phases over two stores"),
        "99700ab1-d770": ("img/fig3c.png", 6.0, "Pairwise F1 against published methods on identical data (Play Store held-out halves; Firefox sample 3)"),
        "20ff9b7c-e968": ("img/fig4c.png", 6.0, "Pairwise precision and recall of every method, Play Store held-out halves"),
        "0e9a5709-bc0f": ("img/fig5c.png", 6.0, "Recall with and without one re-joining pass; precision shown below each pair"),
        "47fef0a8-e40a": ("img/fig6c.png", 6.0, "Share of simulated incidents detected in 6-hour windows; InsightDesk under one false alarm per 100 windows")}

EQS = {
    "6d855770-35fa": r"$\max_{m' \in E_C}\, s(m, m') < 0.85, \qquad |E_C| \leq 5$",
    "cdd1f3c2-8d59": r"$S(x, C) = \max_{m \in E_C}\, s(x, m), \qquad \mathcal{C}_5(x) = \mathrm{top\text{-}5\ cases\ of\ product}(x)$",
    "cce56546-8ef6": [r"$d(x) = \mathrm{join}\ C^*$  if  $S^* \geq t_{high}$",
                      r"$d(x) = J(x, \mathcal{C}_5(x))$  if  $t_{low} \leq S^* < t_{high}$",
                      r"$d(x) = \mathrm{new\ case}$  if  $S^* < t_{low}$"],
    "6ea6c34b-d4df": [r"$\mathrm{Prec}(t) = \dfrac{|\{(i,j): s_{ij} \geq t,\ y_i = y_j\}|}{|\{(i,j): s_{ij} \geq t\}|}$",
                      r"$t_{high} = \min\{\, t : \mathrm{Prec}(t) \geq 0.97 \,\}$"],
    "11929ef3-3264": r"$t_{low} = Q_{0.20}\left(\{\, s_{ij} : y_i = y_j \,\}\right)$",
    "7e2cc8fd-fc63": [r"$t_{low}(x) = 0.40$  if channel$(x) \in$ {app review, chat}",
                      r"$t_{low}(x) = 0.26$  otherwise (tickets, bug reports, web form)"],
    "848a816a-2569": r"$\sigma(A, B) = \max_{a \in A,\ b \in B}\, s(a, b) \geq t_{low}$",
    "6946f64f-abf6": r"$\lambda_{u,t} = \dfrac{\sum_{\tau=t-L}^{t-1} c_{u,\tau} + a}{n_{u,t} + b}, \quad L = 28,\ a = 0.5,\ b = 7$",
    "855df979-22ad": [r"$S_{u,t} = \max(0,\ S_{u,t-1} + c_{u,t} - k_{u,t}), \quad k_{u,t} = \dfrac{\delta}{\ln((\lambda_{u,t} + \delta)/\lambda_{u,t})}$",
                      r"spike if $S_{u,t} \geq h, \quad \delta = 1,\ h = 2$"],
    "d24f008c-2910": r"$\mathrm{alarm}(p,t) \Leftrightarrow v_{p,t} \geq 2 \ \wedge\ P(X \geq v_{p,t}) < 0.001,\quad X \sim \mathrm{Poisson}(\lambda_{p,t})$",
    "37b683ec-8442": r"$\mathrm{eligible}(C) \Leftrightarrow \neg\mathrm{vague}(C) \wedge (n_C \geq 3 \vee \mathrm{alert}(C) \vee \mathrm{alert}(\mathrm{group}(C)))$",
}


def render_eq(ref, n):
    lines = EQS[ref] if isinstance(EQS[ref], list) else [EQS[ref]]
    plt.rcParams.update({"mathtext.fontset": "stix", "font.family": "STIXGeneral"})
    fig = plt.figure(figsize=(6, 0.42 * len(lines) + (0.25 if any("dfrac" in l for l in lines) else 0)), dpi=300)
    for k, line in enumerate(lines):
        fig.text(0.02, 1 - (k + 0.5) / len(lines), line, fontsize=13, va="center")
    path = f"img/eq{n}.png"
    fig.savefig(path, bbox_inches="tight", pad_inches=0.04, transparent=False, facecolor="white")
    plt.close(fig)
    return path


TITLE = ("InsightDesk: LLM-Verified Grouping of Customer Complaints into Problem Cases "
         "for Incident Detection and Escalation")
BODY, SMALL, TAB, REFSZ = 12, 11, 9, 10

# old reference number -> new; recent work (2020-2026) first, older work used as baselines or components last
REMAP = {1: 21, 3: 2, 5: 23, 6: 24, 7: 22, 8: 4, 9: 5, 11: 25, 12: 26, 13: 8, 14: 10, 15: 11, 17: 14, 18: 16,
         19: 27, 20: 28, 21: 17, 22: 19, 24: 20, 25: 29}

REFS = [
    "J. Dąbrowski, E. Letier, A. Perini and A. Susi, “Analysing app reviews for software engineering: a systematic literature review,” Empirical Software Engineering, vol. 27, 2022.",
    "C. Gao et al., “Emerging app issue identification via online joint sentiment-topic tracing,” IEEE Trans. Software Engineering, 2022.",
    "T. Zhang, D. Han, V. Vinayakarao, I. C. Irsan, B. Xu, F. Thung, D. Lo and L. Jiang, “Duplicate bug report detection: How far are we?,” ACM Trans. Software Engineering and Methodology, 2023.",
    "T. Zhang, I. C. Irsan, F. Thung and D. Lo, “Cupid: Leveraging ChatGPT for more accurate duplicate bug report detection,” arXiv:2308.10022, 2023.",
    "R. Peeters, A. Steiner and C. Bizer, “Entity matching using large language models,” in Proc. EDBT, 2025.",
    "L. Zheng et al., “Judging LLM-as-a-judge with MT-Bench and Chatbot Arena,” in Proc. NeurIPS (Datasets and Benchmarks), 2023.",
    "J. Gu et al., “A survey on LLM-as-a-judge,” arXiv:2411.15594, 2024.",
    "A. Abbas, K. Tirumala, D. Simig, S. Ganguli and A. S. Morcos, “SemDeDup: Data-efficient learning at web-scale through semantic deduplication,” arXiv:2303.09540, 2023.",
    "N. Muennighoff, N. Tazi, L. Magne and N. Reimers, “MTEB: Massive text embedding benchmark,” in Proc. EACL, 2023.",
    "J. Liu et al., “Incident-aware duplicate ticket aggregation for cloud systems,” in Proc. ICSE, 2023 (arXiv:2302.09520).",
    "F. Liu et al., “TickIt: Leveraging large language models for automated ticket escalation,” in Proc. FSE (industry), 2025.",
    "T. Ahmed, S. Ghosh, C. Bansal, T. Zimmermann, X. Zhang and S. Rajmohan, “Recommending root-cause and mitigation steps for cloud incidents using large language models,” in Proc. ICSE, 2023.",
    "Y. Chen et al., “Automatic root cause analysis via large language models for cloud incidents,” in Proc. EuroSys, 2024.",
    "SentiSum, “What is anomaly detection?,” product documentation, help.sentisum.com, 2025.",
    "P. Lewis et al., “Retrieval-augmented generation for knowledge-intensive NLP tasks,” in Proc. NeurIPS, 2020.",
    "M. Grootendorst, “BERTopic: Neural topic modeling with a class-based TF-IDF procedure,” arXiv:2203.05794, 2022.",
    "D. Loureiro, F. Barbieri, L. Neves, L. Espinosa Anke and J. Camacho-Collados, “TimeLMs: Diachronic language models from Twitter,” in Proc. ACL (demo), 2022.",
    "F. Barbieri, J. Camacho-Collados, L. Espinosa Anke and L. Neves, “TweetEval: Unified benchmark and comparative evaluation for tweet classification,” in Findings of EMNLP, 2020.",
    "W. Wang, F. Wei, L. Dong, H. Bao, N. Yang and M. Zhou, “MiniLM: Deep self-attention distillation for task-agnostic compression of pre-trained transformers,” in Proc. NeurIPS, 2020.",
    "OpenAI, “gpt-oss-120b and gpt-oss-20b model card,” 2025.",
    "L. Villarroel, G. Bavota, B. Russo, R. Oliveto and M. Di Penta, “Release planning of mobile apps based on user reviews,” in Proc. ICSE, 2016.",
    "C. Sun, D. Lo, S.-C. Khoo and J. Jiang, “Towards more accurate retrieval of duplicate bug reports,” in Proc. ASE, 2011.",
    "N. Reimers and I. Gurevych, “Sentence-BERT: Sentence embeddings using Siamese BERT-networks,” in Proc. EMNLP, 2019.",
    "J. Devlin, M.-W. Chang, K. Lee and K. Toutanova, “BERT: Pre-training of deep bidirectional transformers for language understanding,” in Proc. NAACL, 2019.",
    "G. S. Manku, A. Jain and A. Das Sarma, “Detecting near-duplicates for web crawling,” in Proc. WWW, 2007.",
    "A. Z. Broder, “On the resemblance and containment of documents,” in Proc. Compression and Complexity of Sequences, 1997.",
    "L. McInnes, J. Healy and J. Melville, “UMAP: Uniform manifold approximation and projection for dimension reduction,” arXiv:1802.03426, 2018.",
    "R. J. G. B. Campello, D. Moulavi and J. Sander, “Density-based clustering based on hierarchical density estimates,” in Proc. PAKDD, 2013.",
    "E. S. Page, “Continuous inspection schemes,” Biometrika, vol. 41, pp. 100–115, 1954.",
]

LIT = [
    ("A. App-Review Mining and Emerging-Issue Detection",
     "Mining app reviews for software engineering is a mature field [1]. MERIT [2] tracks topics across app versions "
     "with an online joint sentiment-topic model and flags topics whose share rises in a new version, explaining them "
     "with extracted phrases. Earlier, CLAP [21] grouped reviews that report the same bug with DBSCAN over textual "
     "features for release planning. These systems work per version or per topic rather than per problem case, and "
     "none decides membership online."),
    ("B. Duplicate Detection and LLM Judges",
     "A recent large study of duplicate bug-report detection [3] finds that classic retrieval such as REP [22] remains "
     "competitive with deep models, and Cupid [4] improves it by letting an LLM extract the essentials of each report "
     "first (Recall-Rate@10 0.60–0.65). Bi-encoders such as SBERT [23] give indexable embeddings, while cross-encoders "
     "built on BERT [24] are more accurate but cannot be indexed. LLMs used zero-shot as matchers reach F1 76.4–98.4 on six "
     "entity-matching benchmarks [5], and LLM judges agree with human preferences on many tasks [6], [7]; their known "
     "position and verbosity biases argue for bounding each decision, which InsightDesk does with retrieval and "
     "calibrated thresholds."),
    ("C. Near-Duplicate and Semantic Deduplication",
     "SimHash [25] and MinHash-LSH [26] detect copies at web scale, and SemDeDup [8] removes pairs of embeddings "
     "closer than a small distance ε. Embedding benchmarks [9] show that no single model leads on every clustering "
     "task. These methods target near-copies, not paraphrases of the same problem."),
    ("D. Incident Management and Feedback Analytics",
     "iPACK [10] links customer tickets to incidents using cloud alerts (F1 0.87–0.94), and TickIt [11] groups and "
     "escalates tickets with an LLM (F1 0.88). LLM-based root-cause recommendation for cloud incidents [12], [13] "
     "grounds its explanations in telemetry, in the spirit of retrieval-augmented generation [15]. Commercial "
     "voice-of-customer tools such as SentiSum [14] detect anomalies per tag or theme, typically with z-scores, and "
     "publish neither their grouping method nor its accuracy."),
]

TAB1 = ("COMPARATIVE LITERATURE SURVEY", ["Ref", "Title (venue, year)", "Parameters", "Limitations", "Our Improvement"], [
    ["[2]", "MERIT (TSE, 2022)", "Topic share per app version", "Topics, not problems", "Problem-level cases"],
    ["[3]", "Duplicate bug reports (TOSEM, 2023)", "BM25F / deep retrieval, top-k", "Ranking only, no cases", "Online cases with counts"],
    ["[4]", "Cupid (arXiv, 2023)", "LLM summary + BM25F", "No grouping decision", "Calibrated band + LLM judge"],
    ["[5]", "LLM entity matching (EDBT, 2025)", "Zero-shot prompt per pair", "Pairwise, costly", "One judge call per complaint"],
    ["[8]", "SemDeDup (arXiv, 2023)", "ε = 0.00095–0.03", "Near-copies only", "Paraphrase-level matching"],
    ["[10]", "iPACK (ICSE, 2023)", "Ticket–alert linking", "Needs cloud alerts", "Raw complaint text only"],
    ["[11]", "TickIt (FSE, 2025)", "Predefined categories", "Cannot open new problems", "Open-ended problem cases"],
    ["[13]", "RCACopilot (EuroSys, 2024)", "Telemetry + LLM", "Needs telemetry", "Root cause from complaints"],
    ["[14]", "SentiSum (2025)", "z-score per tag, ~100-day baseline", "Fails at 0–2 per window", "Poisson CUSUM + alert groups"],
])

TAB2 = ("COMPETITORS EVALUATED ON OUR DATA", ["Method", "Parameters", "Limitation", "Our Improvement"], [
    ["CLAP [21]", "eps 0.6 / 0.75", "Shared words merge problems", "LLM judge separates them"],
    ["BM25, REP [22]", "k1 1.2, b 0.75, cut 0.15–0.20", "Misses paraphrases", "Semantic retrieval"],
    ["MinHash-LSH [26]", "3-char shingles, J ≥ 0.15", "Near-copies only", "Matches meaning"],
    ["SimHash [25]", "36–42 of 64 bits", "Near-copies only", "Matches meaning"],
    ["BERTopic [16]", "Library defaults", "Themes, not problems", "Symptom-level cases"],
    ["Embedding cut [23]", "One cosine threshold", "Unsafe or too narrow", "Two cuts + judge"],
    ["Cross-encoder [24]", "Score cut 0.4834", "Fails in a stream", "Stream-calibrated band"],
    ["LLM labellers", "Codebook or own problem list", "Codebook knew the test problems; own list F1 0.32", "Online, no codebook"],
    ["LLM batch grouping", "One prompt per batch", "Order-sensitive (89% of pairs kept); no memory", "Stable cases"],
    ["z-score rule [14]", "z > 2 per tag", "Unstable at small counts", "Poisson CUSUM"],
])

TAB3 = ("RESULTS ON UNSEEN DATA (UPPER THRESHOLD FIXED)", ["Setting", "Data", "Precision", "Note"], [
    ["Second domain", "Firefox, 3 samples", "0.92–0.96", "F1 0.78 with re-joining"],
    ["Second domain", "Thunderbird", "0.87", "F1 0.79 with re-joining"],
    ["Unseen products", "Intercom, Jira SM, ServiceNow", "case 0.92", "No tuning on them"],
    ["Unseen company", "Leave-one-company-out", "case 0.89", "Other companies' thresholds"],
    ["No product field", "889 bug reports, one stream", "0.90", "vs 0.91 with product"],
    ["Retrieval at scale", "233,521 bug reports", "RR@5 0.46–0.56", "BM25 0.40–0.44, REP 0.43–0.44"],
    ["Third domain", "1,000 Endava tickets", "case 0.90", "Cross-encoder 0.66"],
])


def remap(t):
    t = t.replace(" [23]", "")
    return re.sub(r"\[(\d+)\]", lambda m: f"[{REMAP.get(int(m.group(1)), m.group(1))}]", t)


def text_of(xml):
    return html.unescape(re.sub(r"<[^>]+>", "", xml))


def add_runs(par, xml, size=BODY):
    for m in re.finditer(r"<text[^>]*>(.*?)</text>", xml, re.S):
        inner = m.group(1)
        r = par.add_run(remap(text_of(inner)))
        r.bold, r.italic = "<bold>" in inner, "<italic>" in inner
        r.font.size, r.font.name = Pt(size), FONT
    if not re.search(r"<text", xml):
        r = par.add_run(remap(text_of(xml)))
        r.font.size, r.font.name = Pt(size), FONT


def set_cols(section, n):
    cols = section._sectPr.xpath("./w:cols")
    c = cols[0] if cols else OxmlElement("w:cols")
    c.set(qn("w:num"), str(n))
    c.set(qn("w:space"), "360")
    if not cols:
        section._sectPr.append(c)


def new_section(doc, ncols):
    s = doc.add_section(WD_SECTION.CONTINUOUS)
    set_cols(s, ncols)
    return s


def para(doc, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=4):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(0)
    return p


def blocks(xml):
    body = xml[xml.find(">") + 1:xml.rfind("</doc>")]
    depth, start, out = 0, 0, []
    for m in re.finditer(r"<(/?)([a-zA-Z]+)([^>]*?)(/?)>", body):
        close, tag, attrs, selfc = m.groups()
        if close:
            depth -= 1
        elif not selfc:
            depth += 1
        if depth == 0:
            out.append(body[start:m.end()])
            start = m.end()
    return out


def heading(doc, t):
    p = para(doc, WD_ALIGN_PARAGRAPH.CENTER, 4)
    p.paragraph_format.space_before = Pt(10)
    r = p.add_run(t)
    r.font.size = Pt(BODY)


def subheading(doc, t):
    p = para(doc, WD_ALIGN_PARAGRAPH.LEFT, 3)
    p.paragraph_format.space_before = Pt(4)
    r = p.add_run(t)
    r.italic, r.font.size = True, Pt(BODY)


def body_par(doc, t):
    p = para(doc)
    p.paragraph_format.first_line_indent = Inches(0.2)
    r = p.add_run(t)
    r.font.size = Pt(BODY)


def short_table(doc, num, spec):
    cap, head, rows = spec
    for line, sz in ((f"TABLE {num}", TAB), (cap, TAB)):
        p = para(doc, WD_ALIGN_PARAGRAPH.CENTER, 1)
        r = p.add_run(line)
        r.bold, r.font.size = True, Pt(sz)
    table = doc.add_table(rows=len(rows) + 1, cols=len(head))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ri, row in enumerate([head] + rows):
        for ci, c in enumerate(row):
            cp = table.cell(ri, ci).paragraphs[0]
            r = cp.add_run(c)
            r.font.size, r.font.name, r.bold = Pt(TAB), FONT, ri == 0
    para(doc, after=6)


FIG_N = [0]
DELAY = {"ff082679-c090": 2}   # anchor these figures a few paragraphs later so text fills the page first
PENDING = []


def with_caption(path, width_in, text, n):
    """The figure image with its italic caption drawn underneath, so the two float together."""
    from PIL import Image, ImageDraw, ImageFont
    im = Image.open(path).convert("RGB")
    w, h = im.size
    px_per_pt = w / (width_in * 72)
    font = ImageFont.truetype("C:/Windows/Fonts/timesi.ttf", int(SMALL * px_per_pt))
    draw0 = ImageDraw.Draw(im)
    words, lines, cur = text.split(), [], ""
    for wd in words:                                           # wrap the caption to the image width
        trial = (cur + " " + wd).strip()
        if draw0.textlength(trial, font=font) > w * 0.94 and cur:
            lines.append(cur)
            cur = wd
        else:
            cur = trial
    lines.append(cur)
    line_h = int(SMALL * 1.25 * px_per_pt)
    out = Image.new("RGB", (w, h + line_h * len(lines) + int(8 * px_per_pt)), "white")
    out.paste(im, (0, 0))
    d = ImageDraw.Draw(out)
    for k, line in enumerate(lines):
        d.text((w / 2, h + int(4 * px_per_pt) + k * line_h), line, font=font, fill="black", anchor="ma")
    dest = f"img/fig{n}_cap.png"
    out.save(dest)
    return dest


def float_top(run, width_in):
    """Turn the run's inline picture into a floating one at the top of the page, across both columns,
    with text above and below it (IEEE layout): text keeps filling the page instead of leaving a gap."""
    from docx.oxml import parse_xml
    drawing = run._r.find(qn("w:drawing"))
    inline = drawing.find(qn("wp:inline"))
    ext = inline.find(qn("wp:extent"))
    docpr = inline.find(qn("wp:docPr"))
    graphic = inline.find(qn("a:graphic"))
    from lxml import etree
    anchor = parse_xml(
        '<wp:anchor xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
        'distT="0" distB="114300" distL="0" distR="0" simplePos="0" relativeHeight="251659264" behindDoc="0" '
        'locked="0" layoutInCell="1" allowOverlap="0">'
        '<wp:simplePos x="0" y="0"/>'
        '<wp:positionH relativeFrom="margin"><wp:align>center</wp:align></wp:positionH>'
        '<wp:positionV relativeFrom="margin"><wp:align>top</wp:align></wp:positionV>'
        f'<wp:extent cx="{ext.get("cx")}" cy="{ext.get("cy")}"/>'
        '<wp:effectExtent l="0" t="0" r="0" b="0"/><wp:wrapTopAndBottom/>'
        f'<wp:docPr id="{docpr.get("id")}" name="{docpr.get("name")}"/>'
        '<wp:cNvGraphicFramePr/></wp:anchor>')
    anchor.append(etree.fromstring(etree.tostring(graphic)))
    drawing.replace(inline, anchor)


def figure(doc, ref):
    path, width, cap = FIGS[ref]
    FIG_N[0] += 1
    img = with_caption(path, width, f"Fig. {FIG_N[0]}: {cap}", FIG_N[0])
    p = para(doc, WD_ALIGN_PARAGRAPH.CENTER, 0)
    run = p.add_run()
    run.add_picture(img, width=Inches(width))
    float_top(run, width)


def main():
    xml = open("paper.xml", encoding="utf-8").read()
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name, st.font.size = FONT, Pt(BODY)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin = sec.right_margin = Inches(0.7)
    sec.top_margin, sec.bottom_margin = Inches(0.75), Inches(1)
    set_cols(sec, 1)

    bl = blocks(xml)
    eq_n, in_body, skip_to, i = 0, False, None, 0
    while i < len(bl):
        b = bl[i]
        tag = re.match(r"<(\w+)", b).group(1)
        head = re.search(r"heading='(\d)'", b[:120])
        t = text_of(b).strip()
        if skip_to:
            if head and t.startswith(skip_to):
                skip_to = None
            else:
                i += 1
                continue
        if t.startswith("Editor notes"):
            break
        if tag == "paragraph" and head and head.group(1) == "1":
            p = para(doc, WD_ALIGN_PARAGRAPH.CENTER, 4)
            r = p.add_run(TITLE)
            r.bold, r.font.size = True, Pt(20)
        elif not in_body and tag == "paragraph" and not head and not t.startswith(("Abstract", "Keywords")):
            p = para(doc, WD_ALIGN_PARAGRAPH.CENTER, 1)
            if t.startswith("Department"):
                dept, guide = t.split(" Under the guidance")
                r = p.add_run(dept.rstrip(". ")); r.italic = True; r.font.size = Pt(SMALL)
                p2 = para(doc, WD_ALIGN_PARAGRAPH.CENTER, 1)
                r = p2.add_run("Under the guidance" + guide.rstrip(".")); r.italic = True; r.font.size = Pt(SMALL)
            else:
                r = p.add_run(t); r.font.size = Pt(BODY if "Krushnam" in t else SMALL)
            if t.startswith("Academic"):
                new_section(doc, 2)
                in_body = True
        elif t.startswith(("Abstract", "Keywords")):
            p = para(doc, after=6)
            add_runs(p, b, SMALL)
            for r in p.runs:
                r.bold = r.text.strip() in ("Abstract", "Keywords", "—")
                r.italic = True
        elif tag == "paragraph" and head and head.group(1) == "2":
            heading(doc, t.upper())
            if t.startswith("II. LITERATURE"):
                body_par(doc, "Work relevant to InsightDesk falls into four families, compared with their parameters in Table I.")
                for sub, txt in LIT:
                    subheading(doc, sub)
                    body_par(doc, txt)
                skip_to = "E. Research Gaps"
            elif t == "REFERENCES":
                for k, ref in enumerate(REFS, 1):
                    p = para(doc, WD_ALIGN_PARAGRAPH.LEFT, 2)
                    p.paragraph_format.left_indent = Inches(0.3)
                    p.paragraph_format.first_line_indent = Inches(-0.3)
                    r = p.add_run(f"[{k}]  {ref}"); r.font.size = Pt(REFSZ)
                skip_to = "Editor notes"
        elif tag == "paragraph" and head:
            subheading(doc, t)
            if t.startswith("E. Research Gaps"):
                body_par(doc, text_of(bl[i + 1]).strip())
                short_table(doc, "I", TAB1)
                skip_to = "III. PROPOSED"
        elif tag == "paragraph" and re.fullmatch(r"TABLE (II|III)", t):
            short_table(doc, t.split()[1], TAB2 if t.endswith(" II") else TAB3)
            i += 3
            continue
        elif tag == "list":
            for li in re.findall(r"<listItem.*?</listItem>", b, re.S):
                p = para(doc, after=2)
                p.paragraph_format.left_indent = Inches(0.2)
                p.paragraph_format.first_line_indent = Inches(-0.2)
                p.add_run("• ").font.size = Pt(BODY)
                add_runs(p, li, BODY)
        elif tag == "embed":
            ref = re.search(r"ref='node/([^']+)'", b).group(1)
            if ref in FIGS:
                if DELAY.get(ref):
                    PENDING.append([ref, DELAY[ref]])
                else:
                    figure(doc, ref)
            elif ref in EQS:
                eq_n += 1
                from PIL import Image
                img = render_eq(ref, eq_n)
                tab = doc.add_table(rows=1, cols=2)
                tab.autofit = False
                c0, c1 = tab.rows[0].cells
                c0.width, c1.width = Inches(3.0), Inches(0.45)
                p0 = c0.paragraphs[0]
                p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p0.add_run().add_picture(img, width=Inches(min(2.95, Image.open(img).size[0] / 300 * 0.95)))
                p1 = c1.paragraphs[0]
                p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                r = p1.add_run(f"({eq_n})"); r.font.size = Pt(BODY)
        elif tag == "paragraph" and t:
            p = para(doc)
            p.paragraph_format.first_line_indent = Inches(0.2)
            add_runs(p, b, BODY)
            for item in PENDING[:]:
                item[1] -= 1
                if item[1] <= 0:
                    figure(doc, item[0])
                    PENDING.remove(item)
        i += 1
    doc.save("InsightDesk_paper.docx")
    print("saved InsightDesk_paper.docx,", eq_n, "equations")


if __name__ == "__main__":
    main()
